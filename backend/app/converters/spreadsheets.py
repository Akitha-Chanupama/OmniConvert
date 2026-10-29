import csv
import json
from pathlib import Path
from typing import Optional, Set
import openpyxl
from app.converters.base import BaseConverter
from app.schemas.conversion import ConversionOptions

class SpreadsheetConverter(BaseConverter):
    SUPPORTED_INPUTS: Set[str] = {"csv", "xlsx"}
    SUPPORTED_OUTPUTS: Set[str] = {"csv", "xlsx", "json"}

    def can_convert(self, source_ext: str, target_ext: str) -> bool:
        s = source_ext.lower().lstrip(".")
        t = target_ext.lower().lstrip(".")
        valid_pairs = {
            ("csv", "xlsx"),
            ("xlsx", "csv"),
            ("xlsx", "json"),
        }
        return (s, t) in valid_pairs

    def convert(
        self,
        input_path: Path,
        output_path: Path,
        options: Optional[ConversionOptions] = None,
    ) -> Path:
        s = input_path.suffix.lower().lstrip(".")
        t = output_path.suffix.lower().lstrip(".")
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if s == "csv" and t == "xlsx":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            with open(input_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for row_idx, row in enumerate(reader, start=1):
                    for col_idx, val in enumerate(row, start=1):
                        # Try to cast numeric values
                        try:
                            if "." in val:
                                num_val = float(val)
                            else:
                                num_val = int(val)
                            ws.cell(row=row_idx, column=col_idx, value=num_val)
                        except (ValueError, TypeError):
                            ws.cell(row=row_idx, column=col_idx, value=val)

            wb.save(output_path)

        elif s == "xlsx" and t == "csv":
            wb = openpyxl.load_workbook(input_path, data_only=True)
            sheet_name = options.sheet_name if (options and options.sheet_name) else None
            if sheet_name and sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.active

            with open(output_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                for row in ws.iter_rows(values_only=True):
                    # Write row, converting None to empty string
                    writer.writerow(["" if cell is None else str(cell) for cell in row])

        elif s == "xlsx" and t == "json":
            wb = openpyxl.load_workbook(input_path, data_only=True)
            sheet_name = options.sheet_name if (options and options.sheet_name) else None
            ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active

            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                records = []
            else:
                # First row as headers
                headers = [str(c) if c is not None else f"Column_{i}" for i, c in enumerate(rows[0], start=1)]
                records = []
                for row in rows[1:]:
                    rec = {}
                    for i, header in enumerate(headers):
                        val = row[i] if i < len(row) else None
                        rec[header] = "" if val is None else val
                    records.append(rec)

            indent = options.indent if (options and options.indent is not None) else 2
            with open(output_path, "w", encoding="utf-8") as f:
                json.dump(records, f, indent=indent, default=str, ensure_ascii=False)

        if not output_path.exists() or output_path.stat().st_size == 0:
            raise RuntimeError(f"Spreadsheet conversion from {s} to {t} resulted in missing or empty output.")

        return output_path
